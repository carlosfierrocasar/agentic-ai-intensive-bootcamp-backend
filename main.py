import os
import json
from typing import Any, Dict, List, Optional
from datetime import date, datetime

from fastapi import Depends, FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import Column, Integer, String, Date, DateTime, ForeignKey, create_engine, func
from sqlalchemy.orm import Session, declarative_base, sessionmaker, relationship
from openpyxl import load_workbook

try:
    from sqlalchemy import JSON
except Exception:  # pragma: no cover
    JSON = None  # type: ignore


DATABASE_URL = os.getenv("DATABASE_URL")
if not DATABASE_URL:
    raise RuntimeError("DATABASE_URL is not set")

try:
    host_part = DATABASE_URL.split("@", 1)[1].split("/", 1)[0]
except Exception:
    host_part = "<unknown>"
print(f"Using DATABASE_URL host: {host_part}")

engine = create_engine(DATABASE_URL)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


class WeekProgress(BaseModel):
    week: int
    modules_completed: int = 0
    total_modules: int
    assessment_pct: int = 0


def _parse_start_date(v):
    """Accepts None, '', 'YYYY-MM-DD' strings, or date objects; returns date or None."""
    if v is None:
        return None
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        try:
            return date.fromisoformat(s)
        except Exception as e:
            raise ValueError("start_date must be in YYYY-MM-DD format") from e
    raise ValueError("start_date must be a date or YYYY-MM-DD string")


class LearnerCreate(BaseModel):
    name: str
    email: Optional[str] = None
    source_role: str
    target_role: str
    start_week: int = Field(ge=1, le=7)
    start_date: Optional[date] = None

    @field_validator("start_date", mode="before")
    @classmethod
    def validate_start_date(cls, v):
        return _parse_start_date(v)


class LearnerUpdate(BaseModel):
    email: Optional[str] = None
    start_date: Optional[date] = None
    day_checks: Optional[Dict[str, bool]] = None

    @field_validator("start_date", mode="before")
    @classmethod
    def validate_start_date(cls, v):
        return _parse_start_date(v)


class ProgressUpdate(BaseModel):
    items: List[WeekProgress]


class AssessmentWebhook(BaseModel):
    email: str
    week: int = Field(ge=1, le=7)
    track: str
    score: int = Field(ge=0, le=10)


class ExcelSyncInput(BaseModel):
    filename: str = "assessment_results.xlsx"
    week: int = Field(default=1, ge=1, le=7)
    track: str = "engineer"


class LearnerOut(BaseModel):
    id: int
    name: str
    email: Optional[str] = None
    assessment_pct: int = 0
    source_role: str
    target_role: str
    start_week: int
    start_date: Optional[date] = None
    day_checks: Dict[str, bool] = Field(default_factory=dict)
    progress: List[WeekProgress]
    overall_modules_completed: int
    overall_modules_total: int
    overall_progress_pct: int


PASSING_SCORE = 7


def _week_totals() -> List[int]:
    return [5, 5, 5, 5, 5, 5, 4]


def _default_progress() -> List[dict]:
    totals = _week_totals()
    return [
        {"week": i + 1, "modules_completed": 0, "total_modules": totals[i], "assessment_pct": 0}
        for i in range(7)
    ]


def _default_day_checks() -> Dict[str, bool]:
    return {}


def _normalize_progress(progress) -> List[dict]:
    if isinstance(progress, str):
        try:
            progress = json.loads(progress)
        except Exception:
            progress = _default_progress()

    if not isinstance(progress, list):
        progress = _default_progress()

    totals = _week_totals()
    normalized: List[dict] = []
    for i in range(7):
        week_num = i + 1
        base = {
            "week": week_num,
            "modules_completed": 0,
            "total_modules": totals[i],
            "assessment_pct": 0,
        }
        found = next(
            (
                p for p in progress
                if isinstance(p, dict) and int(p.get("week", 0) or 0) == week_num
            ),
            None,
        )
        if isinstance(found, dict):
            base["modules_completed"] = int(found.get("modules_completed", 0) or 0)
            base["total_modules"] = int(found.get("total_modules", totals[i]) or totals[i])
            base["assessment_pct"] = int(found.get("assessment_pct", 0) or 0)

        base["modules_completed"] = max(0, min(base["modules_completed"], base["total_modules"]))
        normalized.append(base)

    return normalized


def _normalize_day_checks(day_checks) -> Dict[str, bool]:
    if isinstance(day_checks, str):
        try:
            day_checks = json.loads(day_checks)
        except Exception:
            day_checks = _default_day_checks()

    if not isinstance(day_checks, dict):
        day_checks = _default_day_checks()

    return {str(k): bool(v) for k, v in day_checks.items()}


def _sync_progress_from_day_checks(progress, day_checks) -> List[dict]:
    normalized = _normalize_progress(progress)
    checks = _normalize_day_checks(day_checks)

    for i, week in enumerate(normalized):
        start_day = i * 5 + 1
        total_days = int(week["total_modules"])
        completed = 0
        for offset in range(total_days):
            if checks.get(f"day_{start_day + offset}", False):
                completed += 1
        week["modules_completed"] = completed

    return normalized


def _sync_day_checks_from_progress(day_checks, progress) -> Dict[str, bool]:
    checks = _normalize_day_checks(day_checks)
    normalized = _normalize_progress(progress)

    for i, week in enumerate(normalized):
        start_day = i * 5 + 1
        total_days = int(week["total_modules"])
        completed = max(0, min(int(week["modules_completed"]), total_days))
        for offset in range(total_days):
            checks[f"day_{start_day + offset}"] = offset < completed

    return checks


def _overall(progress: List[dict]) -> dict:
    total = sum(int(p.get("total_modules", 0)) for p in progress)
    completed = sum(int(p.get("modules_completed", 0)) for p in progress)
    pct = int(round((completed / total) * 100)) if total else 0
    return {
        "overall_modules_total": total,
        "overall_modules_completed": completed,
        "overall_progress_pct": pct,
    }


class Learner(Base):
    __tablename__ = "learners"

    id = Column(Integer, primary_key=True, index=True)
    name = Column(String, nullable=False)
    email = Column(String, nullable=True, index=True)
    assessment_pct = Column(Integer, nullable=False, server_default="0")
    source_role = Column(String, nullable=False)
    target_role = Column(String, nullable=False)
    start_week = Column(Integer, nullable=False)
    start_date = Column(Date, nullable=True)

    if JSON is not None:
        progress = Column(JSON, nullable=False)
        day_checks = Column(JSON, nullable=False, default=dict)
    else:
        from sqlalchemy import Text
        progress = Column(Text, nullable=False)
        day_checks = Column(Text, nullable=False, default="{}")


class Assessment(Base):
    __tablename__ = "assessments"

    id = Column(Integer, primary_key=True, index=True)
    learner_id = Column(Integer, ForeignKey("learners.id"), nullable=False, index=True)
    week = Column(Integer, nullable=False)
    track = Column(String, nullable=False)
    score = Column(Integer, nullable=False)
    submitted_at = Column(DateTime, nullable=False, server_default=func.now())

    learner = relationship("Learner")


app = FastAPI(title="Agentic Bootcamp Backend")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def on_startup():
    Base.metadata.create_all(bind=engine)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def _to_out(row: Learner) -> LearnerOut:
    progress = _normalize_progress(row.progress)
    day_checks = _normalize_day_checks(getattr(row, "day_checks", None))
    progress = _sync_progress_from_day_checks(progress, day_checks)
    ov = _overall(progress)

    return LearnerOut(
        id=row.id,
        name=row.name,
        email=row.email,
        assessment_pct=int(getattr(row, "assessment_pct", 0) or 0),
        source_role=row.source_role,
        target_role=row.target_role,
        start_week=row.start_week,
        start_date=row.start_date,
        day_checks=day_checks,
        progress=[WeekProgress(**p) for p in progress],
        **ov,
    )


def _apply_assessment_result(
    db: Session,
    learner: Learner,
    week: int,
    track: str,
    score: int,
    submitted_at: Optional[datetime] = None,
) -> Assessment:
    assessment = Assessment(
        learner_id=learner.id,
        week=week,
        track=track,
        score=score,
    )
    if submitted_at is not None:
        assessment.submitted_at = submitted_at
    db.add(assessment)

    passed = score >= PASSING_SCORE
    if passed and int(getattr(learner, "assessment_pct", 0) or 0) < 100:
        learner.assessment_pct = 100

    progress = learner.progress
    if isinstance(progress, str):
        try:
            progress = json.loads(progress)
        except Exception:
            progress = _default_progress()
    if not isinstance(progress, list):
        progress = _default_progress()

    for p in progress:
        if isinstance(p, dict) and int(p.get("week", 0) or 0) == week:
            if passed and int(p.get("assessment_pct", 0) or 0) < 100:
                p["assessment_pct"] = 100
            break

    learner.progress = progress
    db.add(learner)
    return assessment


def _parse_excel_datetime(value: Any) -> Optional[datetime]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        for fmt in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%m/%d/%Y %H:%M:%S",
            "%m/%d/%Y %H:%M",
            "%m/%d/%Y",
        ):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                pass
        try:
            return datetime.fromisoformat(text)
        except ValueError:
            return None
    return None


def _coerce_score(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except Exception:
        return None


@app.post("/assessment-webhook")
def assessment_webhook(payload: AssessmentWebhook, db: Session = Depends(get_db)):
    learner = db.query(Learner).filter(Learner.email == payload.email).first()
    if not learner:
        raise HTTPException(status_code=404, detail="Learner not found for email")

    assessment = _apply_assessment_result(
        db=db,
        learner=learner,
        week=payload.week,
        track=payload.track,
        score=payload.score,
    )
    db.commit()
    db.refresh(assessment)

    return {
        "status": "saved",
        "assessment_id": assessment.id,
        "passed": payload.score >= PASSING_SCORE,
    }


@app.post("/sync-assessments-from-excel")
def sync_assessments_from_excel(payload: ExcelSyncInput, db: Session = Depends(get_db)):
    file_path = os.path.join(os.path.dirname(__file__), payload.filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail=f"Excel file not found: {payload.filename}")

    wb = load_workbook(filename=file_path, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise HTTPException(status_code=400, detail="Excel file is empty")

    headers = {str(value).strip(): idx for idx, value in enumerate(header_row) if value is not None}
    required = ["Email", "Score", "Date"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing required Excel columns: {', '.join(missing)}")

    created = 0
    skipped: List[Dict[str, Any]] = []

    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        email_raw = values[headers["Email"]] if headers["Email"] < len(values) else None
        score_raw = values[headers["Score"]] if headers["Score"] < len(values) else None
        date_raw = values[headers["Date"]] if headers["Date"] < len(values) else None

        email = str(email_raw).strip() if email_raw is not None else ""
        score = _coerce_score(score_raw)
        submitted_at = _parse_excel_datetime(date_raw)

        if not email or score is None:
            skipped.append({"row": row_number, "reason": "Missing email or score"})
            continue

        learner = db.query(Learner).filter(Learner.email == email).first()
        if not learner:
            skipped.append({"row": row_number, "email": email, "reason": "Learner not found"})
            continue

        duplicate_query = db.query(Assessment).filter(
            Assessment.learner_id == learner.id,
            Assessment.week == payload.week,
            Assessment.track == payload.track,
            Assessment.score == score,
        )
        if submitted_at is not None:
            duplicate_query = duplicate_query.filter(Assessment.submitted_at == submitted_at)

        if duplicate_query.first() is not None:
            skipped.append({"row": row_number, "email": email, "reason": "Duplicate assessment"})
            continue

        _apply_assessment_result(
            db=db,
            learner=learner,
            week=payload.week,
            track=payload.track,
            score=score,
            submitted_at=submitted_at,
        )
        created += 1

    db.commit()
    return {
        "status": "ok",
        "created": created,
        "skipped": skipped,
    }


@app.get("/learners", response_model=List[LearnerOut])
def list_learners(db: Session = Depends(get_db)):
    rows = db.query(Learner).order_by(Learner.id.asc()).all()
    return [_to_out(r) for r in rows]


@app.post("/learners", response_model=LearnerOut)
def create_learner(payload: LearnerCreate, db: Session = Depends(get_db)):
    row = Learner(
        name=payload.name,
        email=payload.email,
        source_role=payload.source_role,
        target_role=payload.target_role,
        start_week=payload.start_week,
        start_date=payload.start_date,
        progress=_default_progress(),
        day_checks=_default_day_checks(),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return _to_out(row)


@app.patch("/learners/{learner_id}", response_model=LearnerOut)
def update_learner(learner_id: int, payload: LearnerUpdate, db: Session = Depends(get_db)):
    row = db.query(Learner).filter(Learner.id == learner_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Learner not found")

    row.start_date = payload.start_date
    if payload.email is not None:
        row.email = payload.email
    if payload.day_checks is not None:
        normalized_day_checks = _normalize_day_checks(payload.day_checks)
        row.day_checks = normalized_day_checks
        row.progress = _sync_progress_from_day_checks(row.progress, normalized_day_checks)
    db.add(row)
    db.commit()
    db.refresh(row)
    return _to_out(row)


@app.delete("/learners/{learner_id}")
def delete_learner(learner_id: int, db: Session = Depends(get_db)):
    row = db.query(Learner).filter(Learner.id == learner_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Learner not found")
    db.delete(row)
    db.commit()
    return {"status": "deleted"}


@app.put("/learners/{learner_id}/progress", response_model=LearnerOut)
def update_progress(learner_id: int, payload: ProgressUpdate, db: Session = Depends(get_db)):
    row = db.query(Learner).filter(Learner.id == learner_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="Learner not found")

    items = []
    for item in payload.items:
        d = item.model_dump()
        d["week"] = int(d.get("week", 0) or 0)
        d["modules_completed"] = max(0, int(d.get("modules_completed", 0) or 0))
        d["total_modules"] = max(0, int(d.get("total_modules", 0) or 0))
        d["assessment_pct"] = max(0, min(100, int(d.get("assessment_pct", 0) or 0)))
        items.append(d)

    row.progress = items
    row.day_checks = _sync_day_checks_from_progress(getattr(row, "day_checks", None), items)
    db.add(row)
    db.commit()
    db.refresh(row)
    return _to_out(row)
